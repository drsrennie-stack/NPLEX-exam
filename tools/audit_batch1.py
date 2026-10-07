"""Tag Batch 1 with blueprint competency codes and add audit columns.
Usage: python3 audit_batch1.py <batch1.xlsx> <out.xlsx>
"""
import sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# Q# -> (competency code, style, flags, action)
# style: Vignette = a patient case the answer depends on; Recall = no case, or a case the answer ignores
R, V = "Recall", "Vignette"
T = {
 1: ("CV-4", R, "", "Rewrite as a case"),
 2: ("CV-7", R, "", "Rewrite as a case"),
 3: ("CV-3", V, "", "Keep"),
 4: ("CV-8", R, "", "Rewrite as a case (e.g. anterior ST elevation, which artery)"),
 5: ("CV-1", R, "", "Rewrite as a case"),
 6: ("CV-3", V, "", "Keep"),
 7: ("CV-6", R, "", "Rewrite as a case"),
 8: ("CV-5", R, "", "Rewrite as a case"),
 9: ("CV-7", R, "", "Rewrite as a case"),
 10: ("CV-7", R, "", "Rewrite as a case"),
 11: ("CV-18-A2", V, "Tests the same idea as Q27 (concentric hypertrophy from hypertension).", "Keep; replace Q27"),
 12: ("CV-18-B", V, "", "Keep"),
 13: ("CV-6", R, "Formula recall; too easy for board level.", "Replace with a calculation from case data"),
 14: ("CV-13", R, "Patient mentioned, but the answer does not use the case.", "Rewrite so the case matters"),
 15: ("CV-13", R, "", "Rewrite as a case"),
 16: ("CV-18-I4", V, "", "Keep"),
 17: ("CV-18-C3", R, "SEA tagged Biochemistry; this is cell injury, so Pathology.", "Retag SEA; rewrite as a case"),
 18: ("CV-13", V, "", "Keep"),
 19: ("CV-18-D5", R, "Overlaps Q22 (both rheumatic fever).", "Rewrite as a case or replace to widen coverage"),
 20: ("CV-18-K1", R, "Overlaps Q24 (both endocarditis).", "Rewrite as a case"),
 21: ("CV-18-K2", V, "", "Keep"),
 22: ("CV-16", R, "", "Rewrite as a case"),
 23: ("CV-18-K3", R, "", "Rewrite as a case"),
 24: ("CV-18-K1", V, "", "Keep"),
 25: ("CV-18-I3", R, "", "Rewrite as a case"),
 26: ("CV-18-C3", R, "Timing is off. Neutrophils peak at about 1 to 3 days; by 3 to 7 days macrophages are arriving, and granulation tissue forms around 7 to 10 days. The stem's 3 to 5 days straddles the transition and the explanation's day 7 to 10 macrophage peak is not the standard timeline.", "Fix: change stem to 2 days and correct the explanation timeline"),
 27: ("CV-18-A2", V, "Duplicate of Q11.", "Replace with a different hypertensive complication"),
 28: ("CV-18-I10", V, "", "Keep"),
 29: ("CV-18-G4", V, "", "Keep"),
 30: ("CV-11", R, "Tagged Pathology and Level 3, but it is a physiology sequence recall. Option A's arrow order reads as if renin becomes angiotensinogen.", "Retag Physiology Level 1; rewrite as a case (e.g. renal artery stenosis)"),
 31: ("GI-5", R, "", "Rewrite as a case"),
 32: ("GI-13-J1", V, "", "Keep"),
 33: ("GI-2", R, "", "Rewrite as a case"),
 34: ("GI-3", R, "", "Rewrite as a case"),
 35: ("GI-6", R, "", "Rewrite as a case"),
 36: ("GI-13-F5", V, "", "Keep"),
 37: ("GI-6", R, "Overlaps Q38 (both GI hormones).", "Rewrite as a case"),
 38: ("GI-6", R, "", "Rewrite as a case"),
 39: ("GI-12", R, "Tagged Physiology; urease is a virulence factor, so Microbiology & Immunology. Overlaps Q51 (both H. pylori).", "Retag SEA; merge with Q51 into one case"),
 40: ("GI-6", R, "", "Rewrite as a case"),
 41: ("GI-13-E2", V, "", "Keep"),
 42: ("GI-13-L3", V, "", "Keep"),
 43: ("GI-13-K3", V, "", "Keep"),
 44: ("GI-9", R, "", "Rewrite as a case"),
 45: ("GI-13-C3", V, "", "Keep"),
 46: ("GI-13-C2", V, "", "Keep"),
 47: ("GI-8", V, "Hemochromatosis is not named in the blueprint; fits under GI-8 (minerals).", "Keep"),
 48: ("NE-20-E3", R, "The blueprint lists PKU under Neurological, not Gastrointestinal.", "Retag system; rewrite as a case"),
 49: ("GI-13-L7", R, "Negative stem (does NOT). Board items avoid these. Overlaps Q50.", "Rewrite as a positive-stem case"),
 50: ("GI-13-L7", R, "", "Rewrite as a case"),
 51: ("GI-12", R, "Overlaps Q39.", "Merge with Q39"),
 52: ("GI-13-L1", V, "Shiga toxin removes an adenine from 28S rRNA (N-glycosidase); cleaves is loose wording.", "Tighten option B wording"),
 53: ("GI-11", V, "", "Keep"),
 54: ("GI-13-L3", R, "", "Rewrite as a case"),
 55: ("GI-13-G9", R, "", "Rewrite as a case"),
 56: ("GI-13-I1", V, "", "Keep"),
 57: ("GI-13-G2", V, "", "Keep"),
 58: ("GI-13-K6", V, "", "Keep"),
 59: ("GI-13-K7", R, "", "Rewrite as a case"),
 60: ("GI-13-C1", V, "Primary biliary cholangitis is not named in the blueprint; fits under cholestasis.", "Keep"),
}


def main(src, dst):
    wb = load_workbook(src)
    ws = wb["Diagnostic Bank"]
    hdr = [c.value for c in ws[2]]
    start = hdr.index("Status") + 2
    labels = ["Competency Ref", "Style", "Audit notes", "Action", "Correct is longest option"]
    navy = "1E3D4C"
    for i, h in enumerate(labels):
        c = ws.cell(row=2, column=start + i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=navy)
    for row in ws.iter_rows(min_row=3):
        q = row[0].value
        if not isinstance(q, int) or q not in T:
            continue
        code, style, note, act = T[q]
        opts = [len(str(row[13 + i].value or "")) for i in range(4)]
        k = "ABCD".index(row[17].value)
        longest = opts[k] == max(opts) and opts.count(max(opts)) == 1
        for i, v in enumerate([code, style, note, act, "Yes" if longest else ""]):
            c = ws.cell(row=row[0].row, column=start + i, value=v)
            c.alignment = Alignment(vertical="top", wrap_text=(i == 2 or i == 3))
    for i, w in enumerate([12, 10, 60, 34, 12]):
        ws.column_dimensions[ws.cell(row=2, column=start + i).column_letter].width = w
    wb.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
