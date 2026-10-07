"""Build nplex-competencies.js and a Competencies tab in the Master Blueprint Scaffold.

Usage:
  python3 build_competencies.py <scaffold.xlsx> <out_dir>

Code scheme (matches the scaffold legend example CV-18-A3):
  CV-7       numbered competency 7 in Cardiovascular
  CV-18-C    condition category C under competency 18
  CV-18-C3   condition 3 in that category
Spaced recall schedules at the numbered-competency and category level. Questions
may be tagged at the condition level; they roll up to their category.
"""
import json, sys, os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from blueprint_data import EXAM, GEAS, SEAS, SYSTEMS, category_seas

SEA_NAME = {s["code"]: s["name"] for s in SEAS}
SEA_GEA = {s["code"]: s["gea"] for s in SEAS}
GEA_NAME = {g["code"]: g["name"] for g in GEAS}


def dok_for(statement, is_conditions):
    if is_conditions:
        return 3
    return 2


def build():
    comps, flat = [], []  # comps: schedulable units; flat: every code for the sheet
    for s in SYSTEMS:
        for num, short, seas, nd, stmt in s["items"]:
            code = f"{s['code']}-{num}"
            prim = seas[0]
            c = {
                "id": code, "code": code, "system": s["name"], "general": s["name"],
                "systemCode": s["code"], "week": s["week"], "reinforce": s["reinforce"],
                "weight": s["weight"], "kind": "competency", "num": num,
                "name": short, "can": stmt,
                "sea": SEA_NAME[prim], "seaCodes": list(seas),
                "gea": GEA_NAME[SEA_GEA[prim]], "dok": dok_for(stmt, False), "nd": bool(nd),
            }
            comps.append(c)
            flat.append(dict(code=code, level="Competency", parent="", **{k: c[k] for k in ("system", "name", "can", "sea", "gea", "nd", "week")}))
        n = s["condNum"]
        flat.append(dict(code=f"{s['code']}-{n}", level="Competency", parent="", system=s["name"],
                         name=f"{s['name']} conditions", can="Explain pathogenesis and identify etiology, risk factors, complications, and clinical characteristics of the conditions listed under this competency.",
                         sea="Pathology", gea="Disease/Dysfunction", nd=False, week=s["week"]))
        for letter, title, conds in s["categories"]:
            seas, nd = category_seas(title)
            code = f"{s['code']}-{n}-{letter}"
            cl = [{"code": f"{code}{i}", "name": nm} for i, nm in enumerate(conds, 1)]
            stmt = ("Explain pathogenesis and identify etiology, risk factors, complications, and clinical characteristics of "
                    + title[0].lower() + title[1:] + ": " + "; ".join(conds) + ".")
            prim = seas[0]
            c = {
                "id": code, "code": code, "system": s["name"], "general": s["name"],
                "systemCode": s["code"], "week": s["week"], "reinforce": s["reinforce"],
                "weight": s["weight"], "kind": "conditions", "num": n, "letter": letter,
                "name": title, "can": stmt, "conditions": cl,
                "sea": SEA_NAME[prim], "seaCodes": list(seas),
                "gea": GEA_NAME[SEA_GEA[prim]], "dok": 3, "nd": bool(nd),
            }
            comps.append(c)
            flat.append(dict(code=code, level="Category", parent=f"{s['code']}-{n}", system=s["name"], name=title, can=stmt, sea=c["sea"], gea=c["gea"], nd=c["nd"], week=s["week"]))
            for x in cl:
                flat.append(dict(code=x["code"], level="Condition", parent=code, system=s["name"], name=x["name"], can="", sea=c["sea"], gea=c["gea"], nd=c["nd"], week=s["week"]))
    return comps, flat


def blueprint_obj():
    systems = []
    for s in SYSTEMS:
        q = round(s["weight"] * EXAM["items"] / 100)
        systems.append({"code": s["code"], "name": s["name"], "full": s["full"], "weight": s["weight"],
                        "items": q, "cases": q // EXAM["perCluster"], "week": s["week"], "reinforce": s["reinforce"]})
    assert sum(x["items"] for x in systems) == EXAM["items"]
    assert sum(x["cases"] for x in systems) == EXAM["clusters"]
    return {"source": "NABNE NPLEX Part I Biomedical Sciences Study Guide, revised March 2026 (August 2026 administrations)",
            "sourceUrl": "https://www.nabne.org/pdf/NPLEX-Part-I-Biomedical-Sciences-Study-Guide-08-2026.pdf",
            "exam": EXAM, "geas": GEAS, "seas": SEAS, "systems": systems,
            "notes": {"seaTags": "NABNE does not assign SEAs to competencies. seaCodes are our default tags; each question carries its own SEA.",
                      "nd": "nd marks competencies where naturopathic training places extra weight: nutrition biochemistry, antioxidants, detoxification, exercise, and deficiency states. It is our flag, not NABNE's.",
                      "codes": "Codes follow system-number-letter-item, for example CV-18-C3 (myocardial infarction)."}}


def write_js(comps, path):
    head = ("/* ============================================================\n"
            "   NPLEX Part I Biomedical Science, competency map for Mastery OS.\n"
            "   GENERATED by build_competencies.py from the NABNE study guide\n"
            "   (revised March 2026). Do not hand-edit; change blueprint_data.py\n"
            "   and rebuild.\n\n"
            "   Each entry is one spaced-recall unit:\n"
            "     kind 'competency'  a numbered blueprint competency (CV-7)\n"
            "     kind 'conditions'  one condition category (CV-18-C), with its\n"
            "                        conditions listed (CV-18-C3 = MI)\n"
            "   Fields: id, system, week (opens), reinforce (spiral week), weight\n"
            "   (system % of exam), name, can (blueprint statement), sea, seaCodes,\n"
            "   gea, dok (target depth), nd (naturopathic-emphasis flag).\n"
            "   ============================================================ */\n")
    with open(path, "w", encoding="utf-8") as f:
        f.write(head)
        f.write("window.NPLEX_BLUEPRINT = " + json.dumps(blueprint_obj(), ensure_ascii=False, indent=1) + ";\n\n")
        f.write("window.NPLEX_COMPETENCIES = [\n")
        f.write(",\n".join("  " + json.dumps(c, ensure_ascii=False) for c in comps))
        f.write("\n];\n")


def write_sheet(flat, src, dst):
    wb = load_workbook(src)
    if "Competencies" in wb.sheetnames:
        del wb["Competencies"]
    ws = wb.create_sheet("Competencies", 1)
    navy, white = "1E3D4C", "FFFFFF"
    ws["A1"] = "NPLEX Part I Competencies (NABNE study guide, revised March 2026)"
    ws["A1"].font = Font(bold=True, size=14, color=navy)
    ws["A2"] = "Use these codes in the Question Bank's Competency Ref column. Tag at the most specific level that fits (a condition code like CV-18-C3 when the question is about one condition)."
    ws["A2"].alignment = Alignment(wrap_text=False)
    hdr = ["Code", "Level", "Parent", "Body System", "Name", "Blueprint statement", "Default SEA", "GEA", "ND emphasis", "Week"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, color=white)
        c.fill = PatternFill("solid", fgColor=navy)
    for r, x in enumerate(flat, 5):
        vals = [x["code"], x["level"], x["parent"], x["system"], x["name"], x["can"], x["sea"], x["gea"], "Yes" if x["nd"] else "", x["week"]]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.alignment = Alignment(vertical="top", wrap_text=(i == 6))
        if x["level"] != "Condition":
            for i in range(1, 6):
                ws.cell(row=r, column=i).font = Font(bold=True, color=navy)
    widths = [12, 12, 11, 16, 44, 80, 24, 20, 12, 7]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:J{4 + len(flat)}"
    last = 4 + len(flat)

    # Dropdown on Question Bank > Competency Ref
    qb = wb["Question Bank"]
    hdr_row = [c.value for c in qb[2]]
    col = hdr_row.index("Competency Ref") + 1
    letter = qb.cell(row=2, column=col).column_letter
    qb.data_validations.dataValidation = [d for d in qb.data_validations.dataValidation if letter not in str(d.sqref)]
    dv = DataValidation(type="list", formula1=f"=Competencies!$A$5:$A${last}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Unknown code",
                        error="Pick a code from the Competencies tab.")
    dv.add(f"{letter}3:{letter}{qb.max_row}")
    qb.add_data_validation(dv)
    wb.save(dst)


if __name__ == "__main__":
    src, out = sys.argv[1], sys.argv[2]
    os.makedirs(out, exist_ok=True)
    comps, flat = build()
    write_js(comps, os.path.join(out, "nplex-competencies.js"))
    write_sheet(flat, src, os.path.join(out, "NPLEX_Master_Blueprint_Scaffold.xlsx"))
    lv = {}
    for x in flat:
        lv[x["level"]] = lv.get(x["level"], 0) + 1
    print("schedulable units:", len(comps), "| sheet rows by level:", lv)
    for s in SYSTEMS:
        k = [c for c in comps if c["systemCode"] == s["code"]]
        print(f"  {s['code']} {s['name']:16} units {len(k):3}  conditions {sum(len(c.get('conditions', [])) for c in k):3}  nd {sum(c['nd'] for c in k)}")
    seas = {}
    for c in comps:
        seas[c["sea"]] = seas.get(c["sea"], 0) + 1
    print("primary SEA spread:", seas)
